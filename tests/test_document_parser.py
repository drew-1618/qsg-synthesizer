"""Unit tests for multi-format document text extraction."""

import io
import zipfile
from pathlib import Path

import docx
import pytest

from qsg_synthesizer.tools import document_parser
from qsg_synthesizer.tools.document_parser import (
    DocumentParseError,
    UnsupportedDocumentError,
    extension_of,
    extract_text,
    file_kind,
    is_supported,
)

SAMPLE_DOCS_DIR = Path(__file__).parent.parent / "sample_docs"


def _docx_bytes(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestTextFormats:
    def test_markdown_passes_through_cleaned(self):
        text = extract_text("guide.md", b"# Title\n\n\n\nStep  one   here\n")
        assert text == "# Title\n\nStep one here"

    def test_aligned_columns_are_kept_as_tabs(self):
        text = extract_text(
            "guide.txt",
            b"Port    Type\nWAN     RJ-45\n        SFP\n\nOne  stray   gap\n",
        )
        assert text == "Port\tType\nWAN\tRJ-45\n\tSFP\n\nOne stray gap"

    def test_indented_lines_alone_are_not_a_table(self):
        text = extract_text("guide.md", b"- Item\n   - Sub one\n   - Sub two\n")
        assert text == "- Item\n- Sub one\n- Sub two"

    def test_plain_text_strips_bom_and_normalizes_newlines(self):
        raw = "﻿Line one\r\nLine two\rLine three".encode("utf-8")
        assert extract_text("notes.txt", raw) == "Line one\nLine two\nLine three"

    def test_boilerplate_is_stripped_like_pdf_text(self):
        text = extract_text(
            "a.txt", b"Real spec. Page 1 of 4 \xc2\xa9 2024 Acme. All rights reserved."
        )
        assert "Page 1 of 4" not in text
        assert "All rights reserved" not in text
        assert "Real spec." in text

    def test_binary_content_in_a_text_file_is_rejected(self):
        with pytest.raises(DocumentParseError):
            extract_text("fake.txt", b"abc\x00\x01\x02def")

    def test_sample_markdown_matches_what_the_heuristic_expects(self):
        raw = (SAMPLE_DOCS_DIR / "device_a_qsg.md").read_bytes()
        assert "Quick Start Guide" in extract_text("device_a_qsg.md", raw)


class TestDocx:
    def test_headings_lists_paragraphs_and_tables_keep_order(self):
        def build(document):
            document.add_heading("Acme X1 Quick Start", level=1)
            document.add_paragraph("Default IP: 192.168.1.1")
            document.add_paragraph("Plug in power", style="List Bullet")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Port"
            table.cell(0, 1).text = "Speed"
            table.cell(1, 0).text = "WAN | SFP"
            table.cell(1, 1).text = "10G"
            document.add_paragraph("After the table")

        text = extract_text("guide.docx", _docx_bytes(build))
        assert text.index("# Acme X1 Quick Start") < text.index("Default IP")
        assert "- Plug in power" in text
        assert "| Port | Speed |" in text
        assert "| WAN \\| SFP | 10G |" in text
        assert text.index("| Port | Speed |") < text.index("After the table")

    def test_empty_paragraphs_and_tables_are_skipped(self):
        def build(document):
            document.add_paragraph("")
            document.add_table(rows=1, cols=2)
            document.add_paragraph("Only content")

        assert extract_text("a.docx", _docx_bytes(build)) == "Only content"

    def test_non_zip_bytes_are_rejected(self):
        with pytest.raises(DocumentParseError):
            extract_text("a.docx", b"definitely not a zip")

    def test_zip_without_a_word_document_part_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("readme.txt", "hi")
        with pytest.raises(DocumentParseError):
            extract_text("a.docx", buffer.getvalue())

    def test_archives_with_too_many_parts_are_rejected(self, monkeypatch):
        monkeypatch.setattr(document_parser, "_MAX_DOCX_MEMBERS", 2)
        data = _docx_bytes(lambda d: d.add_paragraph("x"))
        with pytest.raises(DocumentParseError, match="too many parts"):
            extract_text("a.docx", data)

    def test_archives_that_expand_too_far_are_rejected(self, monkeypatch):
        monkeypatch.setattr(document_parser, "_MAX_DOCX_UNCOMPRESSED_BYTES", 100)
        data = _docx_bytes(lambda d: d.add_paragraph("x"))
        with pytest.raises(DocumentParseError, match="too much data"):
            extract_text("a.docx", data)


class TestDispatch:
    def test_pdf_goes_through_the_pdf_parser(self):
        raw = (SAMPLE_DOCS_DIR / "device_a_qsg.pdf").read_bytes()
        assert "Quick Start Guide" in extract_text("device_a_qsg.pdf", raw)

    def test_corrupt_pdf_is_a_parse_error(self):
        with pytest.raises(DocumentParseError):
            extract_text("a.pdf", b"not really a pdf")

    def test_unknown_extension_is_unsupported(self):
        with pytest.raises(UnsupportedDocumentError):
            extract_text("deck.pptx", b"x")

    def test_extension_helpers_are_case_insensitive(self):
        assert extension_of("Guide.PDF") == ".pdf"
        assert is_supported("notes.MD")
        assert not is_supported("image.png")
        assert not is_supported("")
        assert file_kind("a.docx") == "Word document"
        assert file_kind("a.pdf") == "PDF"
        assert file_kind("a.md") == "text file"
